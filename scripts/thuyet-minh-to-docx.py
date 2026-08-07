# -*- coding: utf-8 -*-
"""Chuyển bản thuyết minh Markdown sang .docx theo thể thức văn bản hành chính."""
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1]
OUT = sys.argv[2]

FONT = "Times New Roman"
MONO = "Consolas"

doc = Document()

# --- trang A4, lề theo Nghị định 30 ---
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(1.5)


def set_font(run, name=FONT, size=13, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), name)


def base_style():
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(13)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.35
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT)


base_style()

BOLD_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
IMG_RE = re.compile(r"\[ẢNH\s*([0-9]+)\s*:\s*([^\]]+)\]")


def add_rich(par, text, size=13, base_bold=False, base_italic=False):
    """Ghi text có **đậm**, *nghiêng*, `mã` vào một paragraph."""
    for part in BOLD_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size=size, bold=True, italic=base_italic)
        elif part.startswith("`") and part.endswith("`"):
            set_font(par.add_run(part[1:-1]), name=MONO, size=size - 1.5, bold=base_bold)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_font(par.add_run(part[1:-1]), size=size, bold=base_bold, italic=True)
        else:
            set_font(par.add_run(part), size=size, bold=base_bold, italic=base_italic)


def shade(cell, hexcolor):
    tc = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tc.append(el)


TWIPS_PER_CM = 567


def add_page_number_footer(section):
    """Chân trang: 'Trang X' căn giữa."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Trang "), size=11)
    for instr, kind in (("begin", "fldChar"), ("PAGE", "instrText"), ("end", "fldChar")):
        r = p.add_run()
        set_font(r, size=11)
        el = OxmlElement(f"w:{kind}")
        if kind == "fldChar":
            el.set(qn("w:fldCharType"), instr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        r._r.append(el)


def set_row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * TWIPS_PER_CM)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcPr.append(va)


def add_image_box(num, caption):
    """Ô trống chờ dán ảnh + dòng chú thích hình bên dưới."""
    caption = caption[:1].upper() + caption[1:]
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_row_height(t.rows[0], 6.0)
    cell = t.rows[0].cells[0]
    shade(cell, "F7F7F7")
    vcenter(cell)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(f"▣  CHỖ DÁN ẢNH {num}"), size=12, bold=True, color=RGBColor(0xA6, 0x6A, 0x00))

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(2)
    set_font(p2.add_run(caption), size=12, italic=True, color=RGBColor(0x44, 0x44, 0x44))

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(0)
    set_font(
        p3.add_run("(bấm vào ô này → dán ảnh → xoá ba dòng hướng dẫn)"),
        size=10,
        italic=True,
        color=RGBColor(0x90, 0x90, 0x90),
    )

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(12)
    set_font(cap.add_run(f"Hình {num}. {caption}"), size=11, italic=True)


def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        add_rich(p, htxt, size=12, base_bold=True)
        shade(cell, "DCE6F1")
    for row in body:
        cells = t.add_row().cells
        for i, val in enumerate(row[: len(header)]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, val, size=12)
    doc.add_paragraph()


def flush_table(buf):
    if not buf:
        return
    rows = []
    for line in buf:
        if re.match(r"^\|[\s\-:|]+\|$", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if rows:
        add_table(rows)
    buf.clear()


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
table_buf = []
in_code = False
code_buf = []

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    # khối mã / sơ đồ
    if stripped.startswith("```"):
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0)
            for n, cl in enumerate(code_buf):
                r = p.add_run(cl + ("\n" if n < len(code_buf) - 1 else ""))
                set_font(r, name=MONO, size=7.5)
            code_buf = []
            in_code = False
        else:
            flush_table(table_buf)
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line.rstrip())
        i += 1
        continue

    # bảng
    if stripped.startswith("|"):
        table_buf.append(line)
        i += 1
        continue
    flush_table(table_buf)

    if not stripped:
        i += 1
        continue

    # tiêu đề trang bìa
    if stripped.startswith("%%TITLE%%"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(30)
        p.paragraph_format.space_after = Pt(14)
        set_font(p.add_run(stripped.replace("%%TITLE%%", "")), size=20, bold=True)
        i += 1
        continue
    if stripped.startswith("%%SUBTITLE%%"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        add_rich(p, stripped.replace("%%SUBTITLE%%", ""), size=14)
        i += 1
        continue

    # tiêu đề
    if stripped.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        add_rich(p, stripped[5:], size=13, base_bold=True, base_italic=True)
        i += 1
        continue
    if stripped.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        add_rich(p, stripped[4:], size=13.5, base_bold=True)
        i += 1
        continue
    if stripped.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(7)
        add_rich(p, stripped[3:], size=14.5, base_bold=True)
        i += 1
        continue
    if stripped.startswith("# "):
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(14)
        set_font(p.add_run(stripped[2:]), size=16, bold=True, color=RGBColor(0x1F, 0x37, 0x64))
        i += 1
        continue

    # trích dẫn / ghi chú
    if stripped.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.right_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        add_rich(p, stripped[2:], size=13, base_italic=True)
        i += 1
        continue

    # gạch đầu dòng
    if stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, stripped[2:])
        i += 1
        continue

    # danh sách đánh số
    m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
    if m:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, m.group(2))
        i += 1
        continue

    # chỗ chèn ảnh → dựng ô trống chờ dán
    if IMG_RE.search(stripped):
        for num, cap in IMG_RE.findall(stripped):
            add_image_box(num.strip(), cap.strip())
        i += 1
        continue

    # đoạn văn thường
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.6)
    add_rich(p, stripped)
    i += 1

flush_table(table_buf)
add_page_number_footer(sec)
doc.save(OUT)
print("saved:", OUT.encode("ascii", "replace").decode())
